"""Tests for expense workbook loading and profiling."""

from pathlib import Path

import openpyxl
import pytest

from vitali.config import ExpenseDataConfig
from vitali.data.expenses import build_expense_profile, load_expense_data


@pytest.fixture()
def sample_expense_workbooks(tmp_path: Path) -> list[Path]:
    workbook_2025 = tmp_path / "FINCA VITALI _ Contabilidad Administrativa _ 2025.xlsx"
    wb_2025 = openpyxl.Workbook()
    ws_jan = wb_2025.active
    ws_jan.title = "ENERO"
    ws_jan.append(["EGRESOS"])
    ws_jan.append(["FINCA VITALI"])
    ws_jan.append(["Fecha", "Descripción", "Pagado por", "Monto"])
    ws_jan.append(["2025-01-02", "Desinfectante / AMPM", "Pablo", 1295])
    ws_feb = wb_2025.create_sheet("FEBRERO")
    ws_feb.append(["EGRESOS"])
    ws_feb.append(["FINCA VITALI"])
    ws_feb.append(["Fecha", "Descripción", "Pagado por", "Monto"])
    ws_feb.append(["2025-02-10", "Luz / JASEC domos", "Pablo", 43626])
    for month in [
        "MARZO",
        "ABRIL",
        "MAYO",
        "JUNIO",
        "JULIO",
        "AGOSTO",
        "SEPTIEMBRE",
        "OCTUBRE",
        "NOVIEMBRE",
        "DICIEMBRE",
    ]:
        ws = wb_2025.create_sheet(month)
        ws.append(["EGRESOS"])
        ws.append(["FINCA VITALI"])
        ws.append(["Fecha", "Descripción", "Pagado por", "Monto"])
    wb_2025.save(workbook_2025)

    workbook_2026 = tmp_path / "FINCA VITALI _ Contabilidad Administrativa _ 2026.xlsx"
    wb_2026 = openpyxl.Workbook()
    ws_2026 = wb_2026.active
    ws_2026.title = "ENERO"
    ws_2026.append(["EGRESOS"])
    ws_2026.append(["FINCA VITALI"])
    ws_2026.append(["Fecha", "Descripción", "Pagado por", "Monto"])
    ws_2026.append(["2026-01-01", "Starlink / Internet", "Pablo", 23000])
    for month in ["FEBRERO", "MARZO"]:
        ws = wb_2026.create_sheet(month)
        ws.append(["EGRESOS"])
        ws.append(["FINCA VITALI"])
        ws.append(["Fecha", "Descripción", "Pagado por", "Monto"])
    wb_2026.save(workbook_2026)

    return [workbook_2025, workbook_2026]


def test_load_expense_data_reads_rows_from_workbooks(sample_expense_workbooks: list[Path]) -> None:
    config = ExpenseDataConfig(
        source_files=[str(path) for path in sample_expense_workbooks],
    )
    dataset = load_expense_data(sample_expense_workbooks, config)

    assert len(dataset.records) == 3
    assert sorted(dataset.records["source_file"].unique().tolist()) == [
        "FINCA VITALI _ Contabilidad Administrativa _ 2025.xlsx",
        "FINCA VITALI _ Contabilidad Administrativa _ 2026.xlsx",
    ]
    assert dataset.records["amount_crc"].sum() == pytest.approx(1295 + 43626 + 23000)


def test_load_expense_data_rejects_missing_expected_sheet(sample_expense_workbooks: list[Path]) -> None:
    config = ExpenseDataConfig(
        source_files=[str(sample_expense_workbooks[1])],
        expected_sheet_names_by_file={
            "FINCA VITALI _ Contabilidad Administrativa _ 2026.xlsx": ["ENERO", "FEBRERO", "MARZO", "ABRIL"]
        },
    )

    with pytest.raises(ValueError, match="missing expected sheets"):
        load_expense_data([sample_expense_workbooks[1]], config)


def test_build_expense_profile_returns_monthly_summary(sample_expense_workbooks: list[Path]) -> None:
    config = ExpenseDataConfig(source_files=[str(path) for path in sample_expense_workbooks])
    dataset = load_expense_data(sample_expense_workbooks, config)

    profile = build_expense_profile(dataset.records)

    assert profile["source_rows"] == 3
    assert profile["source_files"] == [
        "FINCA VITALI _ Contabilidad Administrativa _ 2025.xlsx",
        "FINCA VITALI _ Contabilidad Administrativa _ 2026.xlsx",
    ]
    assert profile["date_range"] == {
        "expense_date_min": "2025-01-02",
        "expense_date_max": "2026-01-01",
    }
    assert any(row["expense_month"] == 1 for row in profile["by_month"])
