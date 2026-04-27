"""Loading and validation for expense workbooks."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from vitali.config import ExpenseDataConfig

MONTH_TO_NUMBER = {
    "ENERO": 1,
    "FEBRERO": 2,
    "MARZO": 3,
    "ABRIL": 4,
    "MAYO": 5,
    "JUNIO": 6,
    "JULIO": 7,
    "AGOSTO": 8,
    "SEPTIEMBRE": 9,
    "OCTUBRE": 10,
    "NOVIEMBRE": 11,
    "DICIEMBRE": 12,
}


@dataclass
class ExpenseDataset:
    """In-memory representation of expense records parsed from workbooks."""

    records: pd.DataFrame
    source_paths: list[Path]


def load_expense_data(source_files: list[str | Path], config: ExpenseDataConfig) -> ExpenseDataset:
    """Load expense rows from one or more monthly accounting workbooks."""
    source_paths = [Path(path) for path in source_files]
    frames = [_load_expense_workbook(path, config) for path in source_paths]
    records = pd.concat(frames, ignore_index=True) if frames else pd.DataFrame()
    validate_expense_data(records)
    return ExpenseDataset(records=records, source_paths=source_paths)


def _load_expense_workbook(path: Path, config: ExpenseDataConfig) -> pd.DataFrame:
    workbook = load_workbook(path, data_only=True, read_only=True)
    expected_sheets = config.expected_sheet_names_by_file.get(path.name)
    if expected_sheets:
        missing_sheets = sorted(set(expected_sheets) - set(workbook.sheetnames))
        if missing_sheets:
            missing = ", ".join(missing_sheets)
            raise ValueError(f"Workbook {path.name} is missing expected sheets: {missing}")

    rows: list[dict[str, object]] = []
    for sheet_name in workbook.sheetnames:
        worksheet = workbook[sheet_name]
        for raw_row in worksheet.iter_rows(min_row=4, values_only=True):
            date_value, description, paid_by, amount = _extract_expense_fields(raw_row)

            if date_value is None and description is None and amount is None:
                continue

            if not isinstance(amount, (int, float)):
                continue

            rows.append(
                {
                    "source_file": path.name,
                    "source_sheet": sheet_name,
                    "expense_date": pd.to_datetime(date_value, errors="coerce"),
                    "description": str(description).strip() if description is not None else None,
                    "paid_by": str(paid_by).strip() if paid_by is not None else None,
                    "amount_crc": float(amount),
                    "expense_year": _infer_year(path.name, date_value),
                    "expense_month": MONTH_TO_NUMBER.get(sheet_name),
                }
            )

    return pd.DataFrame(rows)


def _extract_expense_fields(raw_row: tuple[object, ...]) -> tuple[object, object, object, object]:
    """Handle both aligned fixtures and real workbooks with a leading blank column."""
    candidates = []
    if len(raw_row) >= 5:
        candidates.append((raw_row[1], raw_row[2], raw_row[3], raw_row[4]))
    if len(raw_row) >= 4:
        candidates.append((raw_row[0], raw_row[1], raw_row[2], raw_row[3]))

    for date_value, description, paid_by, amount in candidates:
        if isinstance(amount, (int, float)):
            return date_value, description, paid_by, amount

    return None, None, None, None


def validate_expense_data(records: pd.DataFrame) -> None:
    """Validate the most important invariants of the parsed expense rows."""
    required_columns = {
        "source_file",
        "source_sheet",
        "expense_date",
        "description",
        "paid_by",
        "amount_crc",
        "expense_year",
        "expense_month",
    }
    missing_columns = sorted(required_columns - set(records.columns))
    if missing_columns:
        missing = ", ".join(missing_columns)
        raise ValueError(f"Expense records are missing required columns: {missing}")

    if records.empty:
        raise ValueError("Expense records are empty.")

    if records["amount_crc"].isna().any():
        raise ValueError("Expense rows must have a numeric amount_crc.")


def build_expense_profile(records: pd.DataFrame) -> dict[str, object]:
    """Produce a compact profile for reproducible expense analysis."""
    by_source_file = []
    for source_file, frame in records.groupby("source_file", dropna=False):
        by_source_file.append(
            {
                "source_file": source_file,
                "rows": int(len(frame)),
                "amount_crc_sum": _to_optional_float(frame["amount_crc"].sum()),
                "sheet_count": int(frame["source_sheet"].nunique()),
            }
        )

    by_month = []
    for (expense_year, expense_month), frame in records.groupby(["expense_year", "expense_month"], dropna=False):
        by_month.append(
            {
                "expense_year": int(expense_year) if pd.notna(expense_year) else None,
                "expense_month": int(expense_month) if pd.notna(expense_month) else None,
                "rows": int(len(frame)),
                "amount_crc_sum": _to_optional_float(frame["amount_crc"].sum()),
            }
        )

    return {
        "source_rows": int(len(records)),
        "source_files": sorted(records["source_file"].dropna().unique().tolist()),
        "date_range": {
            "expense_date_min": _to_optional_date(records["expense_date"].min()),
            "expense_date_max": _to_optional_date(records["expense_date"].max()),
        },
        "by_source_file": by_source_file,
        "by_month": by_month,
        "top_descriptions": [
            {"description": description, "rows": int(count)}
            for description, count in records["description"].value_counts().head(10).items()
        ],
    }


def _infer_year(file_name: str, date_value: object) -> int | None:
    if pd.notna(pd.to_datetime(date_value, errors="coerce")):
        return int(pd.to_datetime(date_value).year)
    for token in ("2025", "2026"):
        if token in file_name:
            return int(token)
    return None


def _to_optional_date(value: object) -> str | None:
    if pd.isna(value):
        return None
    return pd.Timestamp(value).date().isoformat()


def _to_optional_float(value: object) -> float | None:
    if pd.isna(value):
        return None
    return float(value)
